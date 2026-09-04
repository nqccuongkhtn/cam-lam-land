# -*- coding: utf-8 -*-
# BẢNG ĐỔI SỐ TỜ BẢN ĐỒ — 8 XÃ/PHƯỜNG THỊ XÃ NINH HÒA CŨ
# ─────────────────────────────────────────────────────────────────────────────
# NGUỒN: file Excel "Phụ lục 1.10 – 1.17: Số thứ tự các tờ bản đồ" của Văn phòng
# Đăng ký đất đai (chủ dự án gửi 3/9/2026), lưu bản gốc tại
# scripts/dia-chinh/anh-bang-goc/ninh-hoa-phu-luc.xlsx — 8 sheet, mỗi sheet một xã mới.
# Cột C = tên xã mới · D = tờ bản đồ MỚI · E = tên xã cũ · F = tờ bản đồ CŨ.
#
# ⚠️ BẪY 1 — CHỮ TIẾNG VIỆT GÕ TÁCH DẤU: vài sheet lưu chữ kiểu "tách dấu" (NFD),
#    so tên xã kiểu thường sẽ trượt và MẤT TRẮNG khối đầu tiên của sheet đó
#    (đã vấp: đọc lần đầu mất 168 dòng). Phải normalize NFC trước khi so.
#
# ⚠️ BẪY 2 — DỮ LIỆU NINH HÒA KHÁC NHA TRANG. Ở Nha Trang mỗi xã cũ còn một file
#    riêng mang số tờ cũ, nên bảng đổi "tờ mới → (xã cũ, tờ cũ)" dùng được thẳng.
#    Ở Ninh Hòa thì hầu hết xã cũ KHÔNG có file riêng: dữ liệu nằm trong file của
#    xã GỐC, phần lớn là số tờ cũ của chính xã gốc (đổi 1:1), kèm một ít thửa đã
#    được đo vẽ lại và MANG SẴN SỐ TỜ MỚI. Nạp bảng nguyên xi = tra ra 0 kết quả.
#    → Với TỪNG số tờ mới, script tự đếm xem thửa thật nằm ở đâu rồi mới chọn đích:
#         · ở file xã cũ (số tờ cũ)      → ghi đích là (xã cũ, tờ cũ)
#         · ở file xã mới (số tờ mới)    → ghi đích là (xã mới, chính số tờ mới)
#       Chỗ nào cả hai đều không có thửa thì mặc định ghi đích là file xã mới.
#    Khi nào xin được dữ liệu địa chính của các xã cũ thì chạy lại script này.
#
# CHẠY:  python scripts/dia-chinh/bang-so-to-ninh-hoa.py <file.xlsx> <ra.csv>
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from lib_so_to import nfc, num, tao_bang   # phần dùng chung: chuẩn hoá chữ, đổi tên xã → mã, chọn đích, ghi CSV

XLSX = sys.argv[1] if len(sys.argv) > 1 else 'scripts/dia-chinh/anh-bang-goc/ninh-hoa-phu-luc.xlsx'
RA   = sys.argv[2] if len(sys.argv) > 2 else 'scripts/dia-chinh/bang-so-to-ninh-hoa.csv'

# ── 1. Đọc Excel thành các dòng (tên xã mới, tờ mới, tên xã cũ, tờ cũ) ────────
wb = openpyxl.load_workbook(XLSX, data_only=True)
dong = []
for ws in wb.worksheets:
    xaMoi = xaCu = None
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        c, d, e, f = (tuple(r) + (None,) * 6)[2:6]
        if isinstance(c, str) and nfc(c) and 'xã mới' not in nfc(c) and 'Tên ĐVHC' not in nfc(c):
            xaMoi = nfc(c)
        if isinstance(e, str) and nfc(e):
            t = nfc(e)
            if t.lower().startswith(('xã ', 'phường ', 'thị trấn ')) and 'thị xã' not in t.lower():
                xaCu = t
        tm, tc = num(d), num(f)
        if tm is not None and tc is not None and xaMoi and xaCu:
            dong.append((xaMoi, tm, xaCu, tc))
print(f'Đọc được {len(dong)} dòng từ {len(wb.worksheets)} sheet')

tao_bang(dong, RA)
