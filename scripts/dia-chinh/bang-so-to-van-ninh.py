# -*- coding: utf-8 -*-
# BẢNG ĐỔI SỐ TỜ BẢN ĐỒ — 5 XÃ HUYỆN VẠN NINH CŨ
# ─────────────────────────────────────────────────────────────────────────────
# NGUỒN: file Excel "DS Tờ BĐ các xã" của Văn phòng Đăng ký (chủ dự án gửi 4/9/2026),
# lưu bản gốc tại scripts/dia-chinh/anh-bang-goc/van-ninh-to-ban-do.xlsx — MỘT sheet duy nhất,
# xếp theo KHỐI CỘT nằm cạnh nhau (khác hẳn file Ninh Hoà xếp theo sheet):
#     dòng 1 = tên xã MỚI (ô gộp trùm cả nhóm cột)
#     dòng 2 = tên xã CŨ  (ô gộp trùm 2 cột)
#     dòng 3 = "Tờ BĐ cũ" | "Tờ BĐ mới"   ← mỗi xã cũ một cặp cột
#     dòng 4 trở đi = số liệu
# Xã cũ ĐỨNG ĐẦU mỗi nhóm là "đơn vị gốc": cột "Tờ BĐ cũ" BỎ TRỐNG, nghĩa là GIỮ NGUYÊN số tờ.
#
# ĐÃ KIỂM CHỨNG:
#   · Đơn vị gốc khớp 5/5 danh mục của Sở (Vạn Thọ · Vạn Long · Vạn Thắng · Vạn Giã · Vạn Hưng).
#   · Cộng độ dài các khối đúng bằng số tờ cuối cùng của xã mới:
#     Đại Lãnh 129 · Tu Bông 126 · Vạn Thắng 101 · Vạn Ninh 170 · Vạn Hưng 112.
#   · Mọi khối đều liên tục, nối nhau không hở.
#
# CHẠY:  python scripts/dia-chinh/bang-so-to-van-ninh.py <file.xlsx> <ra.csv>
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.utils import column_index_from_string as CI
from lib_so_to import num, tao_bang

XLSX = sys.argv[1] if len(sys.argv) > 1 else 'scripts/dia-chinh/anh-bang-goc/van-ninh-to-ban-do.xlsx'
RA   = sys.argv[2] if len(sys.argv) > 2 else 'scripts/dia-chinh/bang-so-to-van-ninh.csv'

# (tên xã mới, [(tên xã cũ, cột "Tờ BĐ cũ", cột "Tờ BĐ mới"), …]) — đọc từ dòng 1–3 của file.
NHOM = [
    ('xã Đại Lãnh',  [('Vạn Thọ', 'A', 'B'), ('Đại Lãnh', 'C', 'D'), ('Vạn Thạnh', 'E', 'F')]),
    ('xã Tu Bông',   [('Vạn Long', 'H', 'I'), ('Vạn Phước', 'J', 'K'), ('Vạn Khánh', 'L', 'M')]),
    ('xã Vạn Thắng', [('Vạn Thắng', 'O', 'P'), ('Vạn Bình', 'Q', 'R')]),
    ('xã Vạn Ninh',  [('Vạn Giã', 'T', 'U'), ('Vạn Phú', 'V', 'W'), ('Vạn Lương', 'X', 'Y')]),
    ('xã Vạn Hưng',  [('Vạn Hưng', 'AA', 'AB'), ('Xuân Sơn', 'AC', 'AD')]),
]

ws = openpyxl.load_workbook(XLSX, data_only=True).active
hang = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
def o(r, cot):
    i = CI(cot) - 1
    return hang[r - 1][i] if i < len(hang[r - 1]) else None

dong = []
for tenMoi, cap in NHOM:
    cuoi = 0
    for tenCu, cCu, cMoi in cap:
        cu, moi = [], []
        for r in range(4, ws.max_row + 1):
            a, b = num(o(r, cCu)), num(o(r, cMoi))
            if a is not None: cu.append(a)
            if b is not None: moi.append(b)
        # Đơn vị gốc: cột "Tờ BĐ cũ" bỏ trống ⇒ số tờ giữ nguyên.
        if not cu: cu = list(moi)
        if len(cu) != len(moi):
            raise SystemExit(f'{tenMoi} / {tenCu}: cột cũ có {len(cu)} số mà cột mới có {len(moi)} số — lệch nhau.')
        if moi and moi[0] != cuoi + 1:
            raise SystemExit(f'{tenMoi} / {tenCu}: khối bắt đầu ở tờ {moi[0]} nhưng khối trước kết thúc ở {cuoi} — hở/chồng số tờ.')
        for a, b in zip(cu, moi): dong.append((tenMoi, b, tenCu, a))
        cuoi = moi[-1] if moi else cuoi
    print(f'{tenMoi}: {cuoi} tờ')

tao_bang(dong, RA)
