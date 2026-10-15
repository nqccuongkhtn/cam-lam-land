# Tạo MẪU VĂN BẢN đề nghị cung cấp dữ liệu bản đồ địa chính dạng số (file Word) + trang hướng dẫn — 15/10/2026
import os, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf-8')
RA = r'D:\KhoQH\xin-du-lieu\Mau-van-ban-xin-du-lieu-dia-chinh.docx'

d = Document()
for s in d.sections:
    s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(2), Cm(2), Cm(3), Cm(2)
st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(13)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

def doan(chu='', dam=False, nghieng=False, giua=False, co=None, truoc=0, sau=4, thut=None):
    p = d.add_paragraph()
    if giua: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(truoc); p.paragraph_format.space_after = Pt(sau)
    if thut is not None: p.paragraph_format.first_line_indent = Cm(thut)
    if chu:
        r = p.add_run(chu); r.bold = dam; r.italic = nghieng
        if co: r.font.size = Pt(co)
    return p

# ── TRANG 1: VĂN BẢN ĐỀ NGHỊ ──
doan('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', dam=True, giua=True, sau=0)
doan('Độc lập - Tự do - Hạnh phúc', dam=True, giua=True, sau=0)
doan('───────────', giua=True, sau=6)
doan('……………, ngày …… tháng …… năm 2026', nghieng=True, sau=10).alignment = WD_ALIGN_PARAGRAPH.RIGHT
doan('VĂN BẢN ĐỀ NGHỊ', dam=True, giua=True, co=14, sau=0)
doan('Cung cấp thông tin, dữ liệu đất đai (dữ liệu bản đồ địa chính dạng số)', dam=True, giua=True, sau=10)
doan('Kính gửi: Văn phòng Đăng ký đất đai tỉnh ………………………… (hoặc Chi nhánh Văn phòng Đăng ký đất đai ……………………………).', sau=8, thut=1.0)
doan('Căn cứ Luật Đất đai số 31/2024/QH15;', thut=1.0, sau=2)
doan('Căn cứ Nghị định số 101/2024/NĐ-CP ngày 29/7/2024 của Chính phủ quy định về điều tra cơ bản đất đai; đăng ký, cấp Giấy chứng nhận '
     'quyền sử dụng đất, quyền sở hữu tài sản gắn liền với đất và Hệ thống thông tin đất đai (Điều 60 về cung cấp thông tin, dữ liệu đất đai);', thut=1.0, sau=8)
doan('1. Thông tin người đề nghị', dam=True, sau=2)
for dong in ('Tên tổ chức / cá nhân: …………………………………………………………………………',
             'Địa chỉ: ……………………………………………………………………………………………',
             'Mã số thuế / Số định danh cá nhân: ………………………………………………………',
             'Người đại diện: ……………………………………… Chức vụ: ……………………………',
             'Điện thoại: …………………………… Thư điện tử: ……………………………………'):
    doan(dong, thut=1.0, sau=2).alignment = WD_ALIGN_PARAGRAPH.LEFT   # dòng để điền: căn trái, không dãn chữ
doan('2. Nội dung đề nghị cung cấp', dam=True, truoc=6, sau=2)
doan('a) Dữ liệu bản đồ địa chính dạng số (tệp gốc *.dgn hoặc *.shp / *.geojson, hệ tọa độ VN-2000 của tỉnh) của các đơn vị hành chính cấp xã sau:', thut=1.0, sau=4)
bang = d.add_table(rows=5, cols=4); bang.style = 'Table Grid'; bang.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, t in enumerate(('STT', 'Xã / phường (tên hiện hành)', 'Xã / phường trước sắp xếp năm 2025', 'Phạm vi (toàn bộ / các tờ số …)')):
    c = bang.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(t); r.bold = True; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for k in range(1, 5): bang.rows[k].cells[0].text = str(k)
doan('', sau=2)
doan('b) Thuộc tính của từng thửa đất: số tờ bản đồ, số thửa, diện tích, loại đất (mục đích sử dụng đất).', thut=1.0, sau=2)
doan('c) Bảng đối chiếu số hiệu tờ bản đồ địa chính trước và sau khi sắp xếp đơn vị hành chính cấp xã năm 2025 (nếu có).', thut=1.0, sau=2)
doan('Không đề nghị cung cấp thông tin về người sử dụng đất (họ tên, giấy tờ nhân thân, địa chỉ) và các thông tin cá nhân khác.', nghieng=True, thut=1.0, sau=6)
doan('3. Mục đích sử dụng', dam=True, sau=2)
doan('Hiển thị vị trí, ranh giới thửa đất và tra cứu thông tin quy hoạch mang tính tham khảo trên nền tảng bản đồ của đơn vị; không sử dụng vào '
     'mục đích khác; không cung cấp lại dữ liệu gốc cho tổ chức, cá nhân khác.', thut=1.0, sau=6)
doan('4. Hình thức nhận kết quả', dam=True, sau=2)
doan('Tệp dữ liệu qua thư điện tử / thiết bị lưu trữ (USB, đĩa) — nhận trực tiếp tại cơ quan hoặc qua đường bưu chính.', thut=1.0, sau=6)
doan('5. Cam kết', dam=True, sau=2)
doan('Sử dụng thông tin, dữ liệu đúng mục đích đã nêu; nộp phí, chi phí theo quy định; bảo vệ dữ liệu theo quy định của pháp luật '
     '(Nghị định số 13/2023/NĐ-CP về bảo vệ dữ liệu cá nhân); chịu trách nhiệm trước pháp luật về việc sử dụng thông tin, dữ liệu được cung cấp.', thut=1.0, sau=6)
doan('Hồ sơ kèm theo: Phiếu yêu cầu cung cấp thông tin, dữ liệu đất đai (Mẫu số 13/ĐK) và Thông tin, dữ liệu chi tiết về bản đồ địa chính '
     '(Mẫu số 13a/ĐK) ban hành kèm theo Nghị định số 101/2024/NĐ-CP.', nghieng=True, thut=1.0, sau=14)
ky = d.add_table(rows=1, cols=2); ky.alignment = WD_TABLE_ALIGNMENT.CENTER
o = ky.rows[0].cells[1]; o.text = ''
for chu, dam, ngh in (('NGƯỜI ĐỀ NGHỊ', True, False), ('(Ký, ghi rõ họ tên; đóng dấu nếu là tổ chức)', False, True)):
    p = o.add_paragraph() if o.paragraphs[0].text else o.paragraphs[0]
    r = p.add_run(chu); r.bold = dam; r.italic = ngh; p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ── TRANG 2: HƯỚNG DẪN (không gửi kèm) ──
d.add_page_break()
doan('HƯỚNG DẪN SỬ DỤNG MẪU (không gửi kèm văn bản)', dam=True, giua=True, co=14, sau=10)
for chu in (
    '1. Điền đủ mục 1 và bảng ở mục 2; mỗi tỉnh gửi một văn bản riêng tới Văn phòng Đăng ký đất đai của tỉnh đó (dữ liệu bản đồ địa chính do '
    'Văn phòng Đăng ký đất đai / Chi nhánh quản lý).',
    '2. Kèm Phiếu yêu cầu Mẫu số 13/ĐK và Mẫu số 13a/ĐK — tải bản chính thức trong Phụ lục Nghị định số 101/2024/NĐ-CP (Công báo hoặc Cổng thông '
    'tin đất đai quốc gia), đối chiếu bản mới nhất trước khi nộp.',
    '3. Nộp trực tiếp tại bộ phận một cửa, qua bưu chính, hoặc trên Cổng dịch vụ công quốc gia / Cổng dịch vụ công của tỉnh (thủ tục "Cung cấp thông '
    'tin, dữ liệu đất đai").',
    '4. Thông tin có sẵn trong cơ sở dữ liệu được cung cấp trong ngày làm việc (nhận sau 15 giờ thì sang ngày làm việc tiếp theo); nếu từ chối, cơ '
    'quan phải trả lời rõ lý do trong 02 ngày làm việc (theo hướng dẫn thủ tục). Có thể phải nộp phí khai thác dữ liệu đất đai theo quy định.',
    '5. Dữ liệu chỉ được dùng đúng mục đích ghi trong phiếu / văn bản yêu cầu. Thông tin người sử dụng đất là dữ liệu cá nhân — văn bản này đã '
    'ghi rõ KHÔNG xin phần đó.',
    '6. Có tệp về (DGN / shape / GeoJSON) thì chép vào máy, báo cho trợ lý — bộ đổi đã làm cho 387 tờ xã Cam Lâm (scripts/dia-chinh/v8/) dùng lại '
    'được: tự dựng thửa, đối chiếu với web, đưa lên lớp đối chiếu trước khi quyết định thay.',
    '7. Các tỉnh lân cận Khánh Hòa sau sắp xếp năm 2025 (gợi ý): Lâm Đồng (gồm Lâm Đồng, Đắk Nông, Bình Thuận cũ), Đắk Lắk (gồm Đắk Lắk, Phú Yên '
    'cũ), Gia Lai (gồm Gia Lai, Bình Định cũ); phần Ninh Thuận cũ nay thuộc Khánh Hòa — gửi Văn phòng Đăng ký đất đai tỉnh Khánh Hòa.'):
    doan(chu, sau=6)
os.makedirs(os.path.dirname(RA), exist_ok=True)
d.save(RA + '.tmp.docx'); os.replace(RA + '.tmp.docx', RA); print('ghi', RA, os.path.getsize(RA), 'byte')
